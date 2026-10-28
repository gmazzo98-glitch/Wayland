"""AIDA's formatted single-company workbook, separate from row-based exports.

All populated source cells and merged ranges are retained. Financial observations
use their own table's date header and merged-column spans, never column ordinal.
Only explicitly named financial lines feed the existing AIDA formulas.
"""
import hashlib
import io
import json
import math
import re
import unicodedata
from datetime import date, datetime
from pathlib import Path

import openpyxl

DATASET_NAME = "AIDA Single Company Report"
FORMAT = "aida_single_company_report_v1"


def key(value):
    text = unicodedata.normalize("NFKD", str(value or "")).casefold()
    return re.sub(r"[^a-z0-9]+", " ", text).strip()


def safe(value):
    return value.isoformat() if isinstance(value, (datetime, date)) else value


def report_date(value):
    if isinstance(value, (datetime, date)):
        return value.strftime("%Y-%m-%d")
    if isinstance(value, str):
        if "T" in value:
            try:
                return datetime.fromisoformat(value).strftime("%Y-%m-%d")
            except ValueError:
                pass
        for fmt in ("%d/%m/%Y", "%Y-%m-%d"):
            try:
                return datetime.strptime(value.strip(), fmt).strftime("%Y-%m-%d")
            except ValueError:
                pass
    return None


def number(value):
    """Italian printed numbers; unavailable markers are never zero."""
    if isinstance(value, bool) or value is None:
        return None
    if isinstance(value, str):
        value = value.strip().replace("\xa0", "").replace(" ", "")
        if "," in value:
            value = value.replace(".", "").replace(",", ".")
        elif re.fullmatch(r"-?\d{1,3}(?:\.\d{3})+", value):
            value = value.replace(".", "")
    try:
        result = float(value)
        return result if math.isfinite(result) else None
    except (ValueError, TypeError):
        return None


def looks_like_aida_report(frame):
    if frame.empty:
        return False
    tokens = {key(v) for row in frame.head(120).itertuples(index=False, name=None) for v in row if isinstance(v, str)}
    return "codice fiscale" in tokens and bool(tokens & {"anagrafica", "profilo finanziario e dipendenti", "stato patrimoniale"})


_SECTIONS = {
    key(label) for label in (
        "Anagrafica", "Informazioni commerciali e legali", "Informazioni su dimensione e gruppo",
        "Classificazione mercelogica", "Classificazione merceologica", "Gruppo dei pari",
        "Profilo finanziario e dipendenti", "Stato patrimoniale", "Conto economico", "Indici",
        "Probabilit di Default", "Probabilità di Default", "Fondo di Garanzia (DFKA)",
        "Attuali Esponenti/ Manager/ Contatti/ Sindaci", "Consigli e altri organi societari",
        "Management e Staff", "Azionisti controllanti", "Azionisti attuali", "Partecipazioni attuali", "M&A deal",
    )
}

# Source labels, explicitly matched (no fuzzy similarity or inferred financials).
_METRICS = {
    "ricavi delle vendite": "revenue", "a 1 ricavi vendite e prestazioni": "revenue",
    "ebitda": "ebitda", "utile netto": "net_income", "21 utile perdita di esercizio": "net_income",
    "totale attivo": "total_assets", "totale attivit": "total_assets", "totale attivita": "total_assets",
    "dipendenti": "employees", "risultato operativo": "ebit",
    "a tot val della produzione": "production_value", "b costi della produzione": "production_costs",
    "b 6 materie prime e consumo": "materials", "b 9 totale costi del personale": "personnel_costs",
    "c iv tot dispon liquide": "cash", "d totale debiti": "total_debt",
    "b i totale immob immateriali": "intangibles", "a totale patrimonio netto": "equity",
    "patrimonio netto": "equity", "margine sui consumi": "gross_margin",
    "rapporto di indebitamento": "leverage_ratio",
    "grado di copertura degli interessi passivi": "interest_coverage",
}


def parse_report(source, filename=None):
    """Return a JSON-compatible report plus canonical observations and lineage."""
    if isinstance(source, (str, Path)):
        content = Path(source).read_bytes()
        filename = filename or Path(source).name
    else:
        source.seek(0)
        content = source.read()
        filename = filename or getattr(source, "name", "report.xlsx")
    workbook = openpyxl.load_workbook(io.BytesIO(content), data_only=False)
    cached = openpyxl.load_workbook(io.BytesIO(content), data_only=True)
    report = dict(format=FORMAT, filename=filename, sha256=hashlib.sha256(content).hexdigest(),
                  company={}, fields=[], financial_tables=[], people=[], sections=[], sheets=[], warnings=[])
    candidates = {}
    identities = {}
    for ws in workbook.worksheets:
        merges = {(r.min_row, r.min_col): r for r in ws.merged_cells.ranges}
        sheet = dict(name=ws.title, rows=ws.max_row, columns=ws.max_column,
                     merged_ranges=[str(r) for r in ws.merged_cells.ranges], cells=[])
        report["sheets"].append(sheet)
        section = dict(title="Company report", sheet=ws.title, rows=[])
        report["sections"].append(section)
        table, dates, people_header, last_person = None, [], None, None
        for row in ws:
            items = []
            for cell in row:
                if cell.value is None:
                    continue
                merged = merges.get((cell.row, cell.column))
                item = dict(cell=cell.coordinate, value=safe(cell.value),
                            column=cell.column, end_column=merged.max_col if merged else cell.column)
                if cell.data_type == "f":
                    item["formula"] = cell.value
                    item["value"] = safe(cached[ws.title][cell.coordinate].value)
                    if item["value"] is None:
                        report["warnings"].append(f"{ws.title}!{cell.coordinate}: formula has no cached value; kept unavailable.")
                if cell.comment:
                    item["comment"] = cell.comment.text
                if cell.hyperlink:
                    item["hyperlink"] = cell.hyperlink.target
                item["number_format"] = cell.number_format
                items.append(item)
                sheet["cells"].append(item)
            if not items:
                continue
            first = items[0]
            label = key(first["value"])
            if not report["company"].get("legal_name") and len(items) == 1:
                report["company"]["legal_name"] = str(first["value"])
            if len(items) == 1 and label in _SECTIONS:
                section = dict(title=str(first["value"]), sheet=ws.title, rows=[])
                report["sections"].append(section)
                table, dates, people_header, last_person = None, [], None, None
            section["rows"].append(dict(row=first["cell"], cells=items))

            # Explicit label/value pairs, with full addresses for audit.
            if not dates and people_header is None:
                for pos in range(len(items) % 2, len(items) - 1, 2):
                    field = dict(section=section["title"], label=str(items[pos]["value"]),
                                 value=items[pos + 1]["value"], sheet=ws.title,
                                 label_cell=items[pos]["cell"], value_cell=items[pos + 1]["cell"])
                    report["fields"].append(field)
                    if key(field["label"]) in {"codice fiscale", "partita iva"}:
                        identities[str(field["value"])] = field

            date_cells = [(x, report_date(x["value"])) for x in items if report_date(x["value"])]
            if len(date_cells) >= 2 and label.startswith("bilancio"):
                dates = [(x["column"], d) for x, d in date_cells]
                if len(set(d for _, d in dates)) != len(dates):
                    raise ValueError("Duplicate dates in the same financial table; export one accounting basis at a time.")
                table = dict(section=section["title"], sheet=ws.title, basis=str(first["value"]),
                             dates=[d for _, d in dates], units={}, metadata=[], rows=[])
                report["financial_tables"].append(table)
                continue
            if table is not None:
                points = {}
                for col, d in dates:
                    matches = [x for x in items if x["column"] <= col <= x["end_column"]]
                    if len(matches) > 1:
                        raise ValueError(f"Ambiguous financial column at {ws.title}!{first['cell']}.")
                    points[d] = dict(value=matches[0]["value"], cell=matches[0]["cell"]) if matches else dict(value=None, cell=None)
                # Header metadata has no label to the left of the date columns.
                if first["end_column"] >= dates[0][0]:
                    table["metadata"].append(points)
                    for d, p in points.items():
                        if "eur" in key(p["value"]):
                            table["units"][d] = p["value"]
                    continue
                if any(p["cell"] for p in points.values()):
                    metric = dict(label=str(first["value"]).strip(), label_cell=first["cell"], values=points)
                    table["rows"].append(metric)
                    canonical = _METRICS.get(label)
                    if canonical:
                        for d, p in points.items():
                            value = number(p["value"])
                            unit = key(table["units"].get(d, ""))
                            is_ratio = canonical in {"employees", "leverage_ratio", "interest_coverage"}
                            if not is_ratio and unit not in {"migl eur", "eur"}:
                                continue  # preserve unknown units, never mix into k-EUR formulas
                            if value is not None and not is_ratio and unit == "eur":
                                value /= 1000
                            evidence = dict(sheet=ws.title, cell=p["cell"], date=d, label=metric["label"],
                                            raw_value=p["value"], unit=table["units"].get(d), value=value)
                            candidates.setdefault((canonical, d), []).append(evidence)
                    continue

            header_labels = {key(x["value"]): x for x in items}
            if "nome" in header_labels and "ruolo" in header_labels:
                people_header, last_person = header_labels, None
                continue
            if people_header:
                name_col = people_header["nome"]["column"]
                role_col = people_header["ruolo"]["column"]
                role_end = people_header["ruolo"]["end_column"]
                name = next((str(x["value"]) for x in items if x["column"] == name_col), None)
                roles = [str(x["value"]) for x in items if role_col <= x["column"] <= role_end and x["value"] != "-"]
                if name and roles:
                    full_name = re.sub(r"^(?:Sig\.?|Signor[ae]?|Dott\.?|Dott\.ssa)\s+", "", name, flags=re.I)
                    last_person = dict(full_name=full_name, role=" ".join(roles), section=section["title"],
                                       sheet=ws.title, cell=first["cell"], source_rows=[items])
                    report["people"].append(last_person)
                elif last_person:
                    last_person["source_rows"].append(items)
                    if name and re.fullmatch(r"P\d+", name):
                        last_person["person_id"] = name

    if len(identities) != 1:
        raise ValueError("Expected exactly one company fiscal/VAT identifier in this AIDA report.")
    from utils import normalize_registration_nr
    fiscal = next(iter(identities))
    if not re.fullmatch(r"(?:IT)?(?:\d{11}|[A-Z0-9]{16})", fiscal.strip(), re.I):
        raise ValueError("The report's fiscal identifier is not a valid Italian company identifier.")
    report["company"]["registration_number"] = normalize_registration_nr(fiscal, "Italy")
    if not report["financial_tables"] or not report["company"].get("legal_name"):
        raise ValueError("This file does not contain a recognizable AIDA single company report.")

    _company_fields(report)
    unique_people = {}
    for person in report["people"]:
        identity = person.get("person_id") or key(person["full_name"])
        if identity in unique_people:
            existing = unique_people[identity]
            existing["roles"].append(dict(role=person["role"], section=person["section"]))
            existing["role"] += "; " + person["role"]
            existing["source_rows"].extend(person["source_rows"])
        else:
            person["roles"] = [dict(role=person["role"], section=person["section"])]
            unique_people[identity] = person
    report["people"] = list(unique_people.values())
    report["dates"] = sorted({d for t in report["financial_tables"] for d in t["dates"]}, reverse=True)
    latest_year = int(report["dates"][0][:4])
    facts, lineage = {}, {}
    for (metric, d), sources in candidates.items():
        offset = latest_year - int(d[:4])
        suffix = "latest" if offset == 0 else f"y-{offset}"
        name = f"{metric}_{suffix}"
        numbers = {s["value"] for s in sources if s["value"] is not None}
        if len(numbers) > 1 or (name in lineage and lineage[name][0]["date"] != d):
            facts[name] = None
            report["warnings"].append(f"Conflicting observations for {metric} in {d}; excluded from calculated signals.")
        else:
            facts[name] = next(iter(numbers), None)
        lineage.setdefault(name, []).extend(sources)
    facts["accounts_close"] = report["dates"][0]
    facts["derived_last_accounts_year"] = latest_year
    report["facts"], report["lineage"] = facts, lineage
    if facts.get("employees_latest") is not None:
        report["company"]["headcount"] = int(facts["employees_latest"])
    report["cell_count"] = sum(len(s["cells"]) for s in report["sheets"])
    if any("\ufffd" in str(c["value"]) for s in report["sheets"] for c in s["cells"]):
        report["warnings"].append("Some source text already contains damaged characters. Original text has been preserved.")
    return report


def _company_fields(report):
    labels = {key(f["label"]): f["value"] for f in report["fields"]}
    company = report["company"]
    for field, source in (("legal_form", "forma giuridica"), ("registry_status", "stato giuridico"),
                          ("website_url", "sito web"), ("sector_name", "descrizione attivita italiano"),
                          ("sector_name", "descrizione attivit italiano")):
        if source in labels:
            company[field] = str(labels[source])
    company["incorporation_date"] = report_date(labels.get("data di costituzione"))
    ateco = labels.get("ateco 2007") or labels.get("ateco 2025")
    if ateco:
        match = re.match(r"([\d.]+)\s+(.+)", str(ateco))
        if match:
            company["nace_code"] = match[1]
            company.setdefault("sector_name", match[2])
    website = company.get("website_url")
    if website and not website.startswith(("https://", "http://")):
        company["website_url"] = "https://" + website


def signal_drafts(report):
    from aida_import import derive_financial_signals
    return derive_financial_signals(report["facts"])


def plan_import(db, report):
    """Read-only preview. Identity conflicts stop import instead of creating duplicates."""
    from models import Company, RawImportRecord, SignalRecord
    company = db.query(Company).filter_by(registration_number=report["company"]["registration_number"]).first()
    if company is None:
        names = db.query(Company).filter_by(legal_name=report["company"]["legal_name"]).all()
        if names:
            raise ValueError("A company with this name has a different registration number. Resolve its identity before importing.")
    record = db.query(RawImportRecord).filter_by(company_id=company.id, dataset_name=DATASET_NAME).first() if company else None
    signals = {s.signal_key: s for s in db.query(SignalRecord).filter_by(company_id=company.id).all()} if company else {}
    conflicts = []
    for name, draft in signal_drafts(report).items():
        old = signals.get(name)
        if draft["status"] == "present" and old and old.source != DATASET_NAME and old.status != "not_yet_checked" and not old.is_simulated:
            conflicts.append(dict(signal=name, source=old.source, existing_value=old.numeric_value, report_value=draft["value"]))
    changes = [dict(field=k, existing=safe(getattr(company, k)) if company else None, report=v)
               for k, v in report["company"].items() if v is not None and (company is None or safe(getattr(company, k)) != v)]
    return dict(company_id=company.id if company else None, action="Update existing company" if company else "Create company",
                existing_report=record is not None, conflicts=conflicts, changes=changes)


def import_report(db, report, replace_report=False, replace_other_signals=False):
    """One transaction for company, report, people, and source-backed signals."""
    from indicators import fetch_indicator_defs
    from models import Company, CompanyPerson, RawImportRecord, SignalRecord
    plan = plan_import(db, report)
    if plan["existing_report"] and not replace_report:
        raise ValueError("This company already has an AIDA report. Select replace to import again.")
    now = datetime.utcnow()
    try:
        defs = fetch_indicator_defs(db)
        company = db.get(Company, plan["company_id"]) if plan["company_id"] else Company(
            legal_name=report["company"]["legal_name"], registration_number=report["company"]["registration_number"],
            country="Italy", nace_code=None, sector_name=None, shortlist_status="candidate")
        if company.country != "Italy":
            raise ValueError("The matched company's country is not Italy. Resolve its identity before importing.")
        db.add(company)
        for field, value in report["company"].items():
            if value is not None:
                setattr(company, field, datetime.fromisoformat(value) if field == "incorporation_date" else value)
        if company.headcount is not None:
            company.headcount_source_tier = "T1"
            company.segment = "SME" if company.headcount < 250 else "Midcap"
        db.flush()
        existing = {s.signal_key: s for s in db.query(SignalRecord).filter_by(company_id=company.id).all()}
        updated, skipped = 0, []
        from aida_import import FINANCIAL_SIGNAL_INPUT_FIELDS
        for name, draft in signal_drafts(report).items():
            if name not in defs:
                continue
            old = existing.get(name)
            if old and old.source != DATASET_NAME and old.status != "not_yet_checked" and not old.is_simulated and not replace_other_signals:
                if draft["status"] == "present":
                    skipped.append(name)
                continue
            if draft["status"] != "present" and (old is None or old.source != DATASET_NAME):
                continue
            sig = old or SignalRecord(company_id=company.id, signal_key=name)
            sig.source, sig.status, sig.numeric_value = DATASET_NAME, draft["status"], draft["value"]
            sig.text_value, sig.confidence, sig.is_simulated, sig.fetched_at = draft["text"], 1.0, False, now
            fields = [p for k in FINANCIAL_SIGNAL_INPUT_FIELDS.get(name, ()) for p in report["lineage"].get(k, [])]
            sig.raw_payload_ref = json.dumps(dict(dataset=DATASET_NAME, source_file=report["filename"], sha256=report["sha256"],
                                                 basis=draft["basis"], evidence=dict(inputs=draft["inputs"], raw_fields=fields)), ensure_ascii=False)
            db.add(sig)
            updated += 1
        rec = db.query(RawImportRecord).filter_by(company_id=company.id, dataset_name=DATASET_NAME).first()
        if rec is None:
            rec = RawImportRecord(company_id=company.id, dataset_name=DATASET_NAME)
            db.add(rec)
        rec.source_filename, rec.source_row_index, rec.updated_at = report["filename"], 1, now
        rec.raw_row = {**report["facts"], "_aida_report": report}
        rec.mapping_snapshot = {k: f"variable:{k}" for k in report["lineage"]}
        db.query(CompanyPerson).filter_by(company_id=company.id, dataset_name=DATASET_NAME).delete(synchronize_session=False)
        for i, person in enumerate(report["people"]):
            started = re.search(r"dal\s+(\d{2}/\d{2}/\d{4})", person["role"], re.I)
            db.add(CompanyPerson(company_id=company.id, dataset_name=DATASET_NAME, role_group="Flat Import", position_in_row=i,
                                 full_name=person["full_name"], role=person["role"], current_or_former="Current",
                                 appointment_date=datetime.strptime(started[1], "%d/%m/%Y") if started else None,
                                 raw_fields=person, updated_at=now))
        company.need_score, company.readiness_score, company.last_scored_at = None, None, None
        db.commit()
        return dict(company_id=company.id, signals_updated=updated, signals_skipped=skipped, people=len(report["people"]))
    except Exception:
        db.rollback()
        raise
