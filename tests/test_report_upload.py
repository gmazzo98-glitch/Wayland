import json
from datetime import datetime
from io import BytesIO

import openpyxl
import pandas as pd
import pytest
from sqlalchemy import create_engine, event
from sqlalchemy.orm import sessionmaker
from streamlit.testing.v1 import AppTest

from aida_single_report import DATASET_NAME, import_report, parse_report, plan_import
from company_service import parse_uploaded_file
from indicators import seed_indicator_definitions
from models import Base, Company, CompanyPerson, RawImportRecord, SignalRecord


def report_file(missing_middle=False, conflict=False, revenue=6761):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Page 1"
    for address, value in {
        "B1": "EXAMPLE SERVICE S.P.A.", "A3": "80143 Napoli", "T3": "Codice fiscale", "AB3": "06038441215",
        "A5": "Anagrafica", "A7": "Data di costituzione", "M7": datetime(2008, 4, 15),
        "A8": "Ateco 2007", "M8": "649960 Altre intermediazioni finanziarie nca",
        "A9": "80143 Napoli", "T9": "Sito web", "AB9": "www.example.it",
        "A10": "Forma giuridica", "M10": "S.P.A.", "A12": "Profilo finanziario e dipendenti",
    }.items():
        ws[address] = value

    def financial(header, years, rows):
        ws.cell(header, 1, "Bilancio non consolidato")
        for col, year in zip((19, 23, 27, 32), years):
            ws.cell(header, col, f"31/12/{year}")
            ws.cell(header + 1, col, "migl EUR")
        for rownum, label, values in rows:
            ws.cell(rownum, 1, label)
            for first, last, value in zip((17, 21, 25, 30), (20, 24, 29, 34), values):
                ws.merge_cells(start_row=rownum, end_row=rownum, start_column=first, end_column=last)
                if value is not None:
                    ws.cell(rownum, first, value)

    # Unordered dates and header positions inside merged financial value cells.
    financial(14, (2024, 2022, 2023, 2014), [
        (17, "Ricavi delle vendite", (revenue, 5090, None if missing_middle else 6455, 4089)),
        (18, "EBITDA", ("n.d.", "n.d.", "n.d.", "n.d.")),
        (19, "Dipendenti", (18, 20, 17, 16)),
    ])
    ws["A21"] = "Conto economico"
    financial(23, (2023, 2024, 2022, 2014), [
        (26, "A.1. Ricavi vendite e prestazioni", (None if missing_middle else 6455, revenue + int(conflict), 5090, 4089)),
        (27, "21. UTILE/PERDITA DI ESERCIZIO", (259, 204, 25, -10)),
    ])
    for address, value in {
        "A30": "Consigli e altri organi societari", "L32": "Nome", "O32": "Ruolo", "X32": "Organo", "AC32": "Fonte",
        "D33": "1.", "L33": "Sig. Mario Rossi", "O33": "-", "P33": "CONSIGLIERE (dal 28/04/2025 a 31/12/2027)",
        "X33": "BoD", "AC33": "HO", "L34": "P123", "O35": "Cerved (received on 18/08/2026)",
        "A37": "Management e Staff", "L39": "Nome", "O39": "Ruolo", "X39": "Dpt.", "AC39": "Fonte",
        "D40": "1.", "L40": "Sig. Mario Rossi", "O40": "-", "P40": "AMMINISTRATORE DELEGATO",
        "X40": "SenMan", "AC40": "HO", "L41": "P123", "A43": "Azionisti controllanti",
        "C45": "This entity is independent and has no controlling shareholders.",
        "A47": "Probabilità di Default", "A49": "You have no access to Probability of default module.",
    }.items():
        ws[address] = value
    ws.merge_cells("O32:W32")
    ws.merge_cells("O39:W39")
    ws2 = wb.create_sheet("Additional notes")
    ws2["D5"] = "Unrecognised source content must be retained."
    out = BytesIO()
    wb.save(out)
    out.seek(0)
    return out


@pytest.fixture
def db():
    engine = create_engine("sqlite:///:memory:")
    Base.metadata.create_all(engine)
    with sessionmaker(bind=engine)() as session:
        seed_indicator_definitions(session)
        yield session
    engine.dispose()


def test_dates_come_from_each_table_and_missing_cells_never_shift_years():
    report = parse_report(report_file(missing_middle=True))
    assert report["facts"]["revenue_latest"] == 6761
    assert report["facts"]["revenue_y-2"] == 5090
    assert report["facts"]["revenue_y-1"] is None
    assert report["facts"]["revenue_y-10"] == 4089
    assert "revenue_y-9" not in report["facts"]
    assert report["facts"]["ebitda_latest"] is None
    assert report["lineage"]["revenue_y-2"][0]["cell"] == "U17"
    assert report["company"]["registration_number"] == "IT06038441215"
    assert report["company"]["nace_code"] == "649960"
    assert report["company"]["incorporation_date"] == "2008-04-15"
    assert report["company"]["website_url"] == "https://www.example.it"


def test_every_populated_cell_on_every_sheet_survives_with_ownership_and_people():
    source = report_file()
    original = openpyxl.load_workbook(source)
    source.seek(0)
    report = parse_report(source)
    actual = {(s["name"], c["cell"]): c["value"] for s in report["sheets"] for c in s["cells"]}
    expected = {(s.title, c.coordinate): c.value.isoformat() if isinstance(c.value, datetime) else c.value
                for s in original for row in s for c in row if c.value is not None}
    assert actual == expected
    assert report["cell_count"] == len(expected)
    assert len(report["people"]) == 1
    person = report["people"][0]
    assert person["full_name"] == "Mario Rossi" and person["person_id"] == "P123"
    assert len(person["roles"]) == 2
    assert "CONSIGLIERE" in person["role"] and "AMMINISTRATORE DELEGATO" in person["role"]
    assert any(s["title"] == "Azionisti controllanti" for s in report["sections"])
    assert actual[("Page 1", "A49")].startswith("You have no access")


def test_disagreeing_duplicate_financial_observations_are_excluded():
    report = parse_report(report_file(conflict=True))
    assert report["facts"]["revenue_latest"] is None
    assert any("Conflicting observations" in w for w in report["warnings"])


def test_flexible_upload_routes_report_to_dedicated_flow():
    with pytest.raises(ValueError, match="AIDA single company report upload"):
        parse_uploaded_file(report_file(), filename="company.xlsx")


def test_normal_tabular_upload_stays_one_row_per_record():
    path = BytesIO()
    pd.DataFrame([
        {"Legal Name": "Example S.p.A.", "Registration Number": "IT123"},
        {"Legal Name": "Another S.r.l.", "Registration Number": "IT456"},
    ]).to_excel(path, index=False)
    path.seek(0)

    result = parse_uploaded_file(path, filename="table.xlsx")

    assert result.shape == (2, 2)
    assert list(result.columns) == ["legal_name", "registration_number"]


def test_preview_is_read_only_and_confirm_saves_report_people_and_auditable_signals(db):
    report = parse_report(report_file())
    plan = plan_import(db, report)
    assert plan["company_id"] is None
    assert db.query(Company).count() == 0
    result = import_report(db, report)
    company = db.get(Company, result["company_id"])
    assert company.headcount == 18
    assert company.country == "Italy" and company.nace_code == "649960"
    assert db.query(CompanyPerson).count() == 1
    record = db.query(RawImportRecord).one()
    assert isinstance(record.raw_row["_aida_report"], dict)
    assert record.raw_row["_aida_report"]["cell_count"] == report["cell_count"]
    signal = db.query(SignalRecord).filter_by(signal_key="revenue_trend").one()
    assert signal.numeric_value == pytest.approx((6761 - 5090) / 5090 * 100, abs=0.0001)
    payload = json.loads(signal.raw_payload_ref)
    assert payload["sha256"] == report["sha256"]
    assert payload["evidence"]["raw_fields"][0]["sheet"] == "Page 1"
    assert db.query(SignalRecord).filter_by(signal_key="ebitda_trend", status="present").count() == 0


def test_reimport_requires_replace_then_updates_without_duplicate_people(db):
    report = parse_report(report_file())
    result = import_report(db, report)
    with pytest.raises(ValueError, match="already has"):
        import_report(db, report)
    revised = parse_report(report_file(revenue=7000))
    import_report(db, revised, replace_report=True)
    assert db.query(Company).count() == 1
    assert db.query(CompanyPerson).count() == 1
    assert db.query(RawImportRecord).one().raw_row["revenue_latest"] == 7000
    assert db.get(Company, result["company_id"]).headcount == 18


def test_existing_other_source_signal_is_preserved_unless_reviewed(db):
    company = Company(legal_name="EXAMPLE SERVICE S.P.A.", registration_number="IT06038441215", country="Italy")
    db.add(company)
    db.flush()
    db.add(SignalRecord(company_id=company.id, signal_key="revenue_trend", numeric_value=7, source="Manual Entry",
                        status="present", confidence=1, is_simulated=False))
    db.commit()
    report = parse_report(report_file())
    assert len(plan_import(db, report)["conflicts"]) == 1
    import_report(db, report)
    assert db.query(SignalRecord).filter_by(signal_key="revenue_trend").one().numeric_value == 7
    import_report(db, report, replace_report=True, replace_other_signals=True)
    assert db.query(SignalRecord).filter_by(signal_key="revenue_trend").one().source == DATASET_NAME


def test_import_failure_rolls_back_company_and_all_children(db):
    report = parse_report(report_file())

    def reject_people(mapper, connection, target):
        raise RuntimeError("injected failure")

    event.listen(CompanyPerson, "before_insert", reject_people)
    try:
        with pytest.raises(RuntimeError, match="injected failure"):
            import_report(db, report)
    finally:
        event.remove(CompanyPerson, "before_insert", reject_people)
    assert db.query(Company).count() == 0
    assert db.query(RawImportRecord).count() == 0
    assert db.query(SignalRecord).count() == 0


def test_name_collision_never_creates_duplicate_company(db):
    db.add(Company(legal_name="EXAMPLE SERVICE S.P.A.", registration_number="IT11111111111", country="Italy"))
    db.commit()
    with pytest.raises(ValueError, match="different registration number"):
        plan_import(db, parse_report(report_file()))


def test_report_view_renders_tables_and_has_no_mapping_dropdowns():
    report = parse_report(report_file())
    app = AppTest.from_string("""
import streamlit as st
from views.aida_single_report import render_report
render_report(st.session_state['report'])
""")
    app.session_state["report"] = report
    app.run(timeout=30)
    assert not app.exception
    assert len(app.selectbox) == 0
    assert len(app.dataframe) >= 5
    assert any(e.label == "People and appointments" for e in app.expander)


def test_upload_preview_and_confirmation_save_only_on_click():
    app = AppTest.from_string("""
import streamlit as st
from sqlalchemy import create_engine
from sqlalchemy.orm import Session
from sqlalchemy.pool import StaticPool
from unittest.mock import patch
from models import Base, Company, RawImportRecord
from indicators import seed_indicator_definitions
from views.aida_single_report import render_upload
if 'report_db' not in st.session_state:
    engine = create_engine('sqlite:///:memory:', connect_args={'check_same_thread': False}, poolclass=StaticPool)
    Base.metadata.create_all(engine)
    db = Session(engine)
    seed_indicator_definitions(db)
    st.session_state['report_db'] = db
db = st.session_state['report_db']
with patch.object(st, 'file_uploader', return_value=st.session_state['upload']):
    render_upload(db)
st.session_state['saved_companies'] = db.query(Company).count()
st.session_state['saved_reports'] = db.query(RawImportRecord).count()
""")
    source = report_file()
    source.name = "single_company.xlsx"
    app.session_state["upload"] = source
    app.run(timeout=30)
    assert not app.exception
    assert app.session_state["saved_companies"] == 0
    app.button(key="aida_single_confirm").click().run(timeout=30)
    assert not app.exception
    assert app.session_state["saved_companies"] == 1
    assert app.session_state["saved_reports"] == 1
    assert app.success


def test_dossier_financial_statement_reads_the_saved_single_report(db, monkeypatch):
    from types import SimpleNamespace
    import company_service
    from views.company_dossier import _load_shared
    monkeypatch.setattr(company_service, "get_company_pain_points", lambda *args: [])
    result = import_report(db, parse_report(report_file()))
    company = db.get(Company, result["company_id"])
    shared = _load_shared(SimpleNamespace(db=db, company=company, sig_dict={}))
    assert shared["fin"]["labels"]["latest"] == "FY2024"
    assert shared["fin"]["lines"]["revenue"]["latest"] == 6761
    assert shared["fin"]["lines"]["net_income"]["latest"] == 204
